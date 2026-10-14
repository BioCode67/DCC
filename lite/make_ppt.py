"""결과(results.json) → 제출용 PPT 생성 (python-pptx)."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

NAVY = RGBColor(0x1F, 0x3A, 0x68)
GRAY = RGBColor(0x55, 0x55, 0x55)


def _title(slide, text, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text, p.font.size, p.font.bold, p.font.color.rgb = text, Pt(30), True, NAVY
    if sub:
        q = tb.text_frame.add_paragraph()
        q.text, q.font.size, q.font.color.rgb = sub, Pt(16), GRAY


def _bullets(slide, lines, top=1.4, left=0.6, width=12.1, size=18):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(5.8))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        lvl = 1 if line.startswith("  ") else 0
        p.text = ("· " if lvl else "■ ") + line.strip()
        p.level = lvl
        p.font.size = Pt(size - 3 * lvl)
        p.space_after = Pt(6)


def _table(slide, rows, top=1.5, left=0.6, width=12.1, size=14):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width),
                                   Inches(0.45 * len(rows)))
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = shape.table.cell(r, c)
            cell.text = str(v)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(size)
                p.font.bold = r == 0


def _bar(slide, cats, vals, title, left, top, w=6, h=4):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(title, vals)
    ch = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(left), Inches(top),
                                Inches(w), Inches(h), cd).chart
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    ch.plots[0].has_data_labels = True


def _fmt(v, nd=4):
    return f"{v:.{nd}f}" if isinstance(v, (int, float)) else "-"


def build(R, path, team="200OK"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    st = R.get("stats", {})
    tr, va = st.get("train", {}), st.get("val", {})
    env = R.get("env", {})
    m = {n: R.get(f"mission{n}", {}) for n in (1, 2, 3)}

    s = prs.slides.add_slide(blank)
    _title(s, "2026 데이터·AI 크리에이터 캠프 대학부 예선", f"팀 {team} | 전화대화 음성 및 전사 데이터를 이용한 응급상황 인식")
    _bullets(s, ["Mission 1  신고자의 성별 분류 (음성)", "Mission 2  신고자 / 119대원 발화 분류 (음성)",
                 "Mission 3  환자의 증상 인식 (대화 텍스트, 9개 증상 multi-label)"], top=2.6, size=22)

    s = prs.slides.add_slide(blank)
    _title(s, "데이터 분석", "AI-Hub 위급상황 음성/음향(고도화) - 119 지능형 신고접수 음성 인식 데이터 (서울, 구급)")
    _table(s, [["", "Training", "Validation"],
               ["통화 수", tr.get("n_calls", "-"), va.get("n_calls", "-")],
               ["평균 통화 길이(초)", _fmt(tr.get("avg_dur_sec"), 1), _fmt(va.get("avg_dur_sec"), 1)],
               ["통화당 평균 발화 수", _fmt(tr.get("avg_utts"), 1), _fmt(va.get("avg_utts"), 1)],
               ["성별 분포", tr.get("gender", "-"), va.get("gender", "-")],
               ["발화 화자 분포 (0=대원, 1=신고자)", tr.get("speaker", "-"), va.get("speaker", "-")]],
           size=13)
    if tr.get("symptom"):
        s = prs.slides.add_slide(blank)
        _title(s, "데이터 분석 — 증상 분포 (Training)", "9개 대상 증상만 집계, 클래스 불균형이 큼 → Macro F1 에 불리")
        _bar(s, list(tr["symptom"]), list(tr["symptom"].values()), "증상별 통화 수", 0.5, 1.5, 7.4, 5.3)
        ns = tr.get("n_symptom", {})
        _bar(s, [f"{k}개" for k in ns], list(ns.values()), "통화당 증상 개수", 8.1, 1.5, 4.8, 5.3)

    s = prs.slides.add_slide(blank)
    _title(s, "공통 원칙 및 규정 준수")
    _bullets(s, [
        "AI-Hub 제공 데이터만 사용, Training 은 서울 데이터만, Validation 은 학습·모델 선택에 미사용",
        "모델 선택은 Training 내부 10% holdout 으로 수행 (같은 통화는 같은 쪽에 배정해 누수 방지)",
        "미션별 지정 입력만 사용",
        "  M1: 음성 + startAt / endAt / speaker   M2: 조각 경계(startAt / endAt)만, 발화 순서·text 미사용",
        "  M3: 전사 텍스트만 (speaker 등 다른 라벨 미사용)",
        "Decision Threshold 0.5 고정 — 모든 미션·클래스 공통, threshold tuning 없음",
        "label 별로 다른 전처리 없음, 평가 데이터 전용 전처리 없음, 상용 AI API 미사용",
        "inference.py 한 번 실행으로 결과 CSV 생성, 인터넷 불필요",
    ])

    s = prs.slides.add_slide(blank)
    _title(s, "Mission 1 — 신고자 성별 분류", "입력: 음성 + startAt/endAt/speaker → 출력: gender | 평가: Accuracy")
    _bullets(s, [
        "신고자(speaker=1) 발화 구간만 이어서 최대 30초 사용 → 119대원 목소리 배제",
        "16kHz 변환 → 10ms 프레임 단위 MFCC 20차, 기본주파수(F0, YIN), RMS 에너지",
        "구간 통계(평균·표준편차·변화량, F0 분위수) 70차원 → HistGradientBoosting 분류기",
        "근거: 성별은 목소리 높낮이(F0)와 성도 특성(MFCC)에 강하게 반영됨",
        f"결과: Validation Accuracy {_fmt(m[1].get('score'))}  (Training holdout {_fmt(m[1].get('holdout'))})",
    ])

    s = prs.slides.add_slide(blank)
    _title(s, "Mission 2 — 신고자 / 119대원 발화 분류", "입력: 한 발화 조각 → 출력: speaker (1 신고자 / 0 대원) | 평가: Accuracy")
    _bullets(s, [
        "통화 전체 프레임 특징을 한 번 계산하고 각 조각 구간(startAt~endAt)만 잘라 통계 특징 생성",
        "발화 순서·홀짝·text 등은 사용하지 않음, 각 조각을 독립적으로 분류",
        "특징: MFCC 통계 + F0 분위수 + 에너지 + 조각 길이 → HistGradientBoosting",
        "학습: 통화당 최대 10개 조각을 화자별 균형 샘플링",
        "근거: 119대원은 일정한 말투·녹음 채널, 신고자는 긴박한 발화(높은 F0·에너지 변동)",
        f"결과: Validation Accuracy {_fmt(m[2].get('score'))}  (Training holdout {_fmt(m[2].get('holdout'))})",
    ])

    s = prs.slides.add_slide(blank)
    _title(s, "Mission 3 — 환자의 증상 인식", "입력: 신고자·대원 대화 텍스트 → 출력: 9개 증상 (0~9개) | 평가: Macro F1")
    _bullets(s, [
        "발화 텍스트를 순서대로 연결 → 문자 n-gram(2~4) TF-IDF (띄어쓰기·구어체 변형에 강함)",
        "증상별 Logistic Regression 9개 (class_weight='balanced' 로 희소 증상 학습 보완)",
        "9개 대상 외 증상(예: 찰과상)은 라벨에서 제외, 확률 ≥ 0.5 인 증상만 출력",
        f"결과: Validation Macro F1 {_fmt(m[3].get('score'))}  (Training holdout {_fmt(m[3].get('holdout'))})",
    ], size=17)
    per = m[3].get("per_class")
    if per:
        _bar(s, list(per), [round(v, 3) for v in per.values()], "증상별 Validation F1", 0.6, 3.6, 12, 3.6)

    s = prs.slides.add_slide(blank)
    _title(s, "결과 요약 및 계산 효율", f"환경: {env.get('cpu', '-')} ({env.get('cores', '-')} cores), GPU 미사용 | Python {env.get('python', '-')}")
    rows = [["미션", "지표", "Validation", "Train holdout", "Val 추론 시간(전체)", "모델 크기"]]
    for n in (1, 2, 3):
        r = m[n]
        rows.append([f"Mission {n}", r.get("metric", "-"), _fmt(r.get("score")), _fmt(r.get("holdout")),
                     f"{r['infer_sec']:.0f}s ({r.get('n', 0)}건)" if r.get("infer_sec") else "-",
                     f"{r['ckpt_mb']:.1f} MB" if r.get("ckpt_mb") else "-"])
    _table(s, rows, size=14)
    _bullets(s, ["모든 모델은 CPU 만으로 학습·추론 (딥러닝 파라미터 없음, 트리/선형 모델)",
                 "추론 batch: 통화 단위 병렬 처리 (CPU 코어 수 - 1 프로세스)"], top=4.2, size=15)

    s = prs.slides.add_slide(blank)
    _title(s, "한계 및 향후 계획")
    _bullets(s, [
        "컴퓨팅 자원 제약: 무료 Colab 의 시간 제한·Drive I/O 지연으로 GPU 학습을 완료하지 못함",
        "  → 제출은 CPU 경량 모델, 미션 1·2 는 Training 일부 통화만 샘플링해 학습",
        "준비한 딥러닝 버전 (코드 완성, 학습은 향후 과제)",
        "  M1·M2: 사전학습 음성모델 WavLM(base-plus) 특징 + MLP (layer 가중합)",
        "  M3: KLUE-RoBERTa fine-tuning + 긴 대화 window max-pooling",
        "짧은 발화(\"네\", \"예\")는 화자 정보가 부족해 M2 오분류의 주원인으로 예상",
        "유사 증상(구토/오심, 어지러움/전신쇠약) 구분과 희소 증상 재현율 개선 필요",
    ])

    s = prs.slides.add_slide(blank)
    _title(s, "사회안전 시사점")
    _bullets(s, [
        "신고 접수 자동 보조: 통화 중 실시간으로 신고자 발화와 증상을 구조화 → 상황판·출동 지령 입력 시간 단축",
        "화자 분리(M2)는 상담 기록 자동화와 신고자 음성만의 감정·긴급도 분석의 전처리로 활용",
        "증상 인식(M3)은 중증도 분류(triage)와 병원 선정, 응급처치 지도(CPR 등) 안내를 앞당김",
        "성별 등 신고자 특성(M1)은 신원 확인이 어려운 상황에서 현장 대원에게 보조 정보 제공",
        "경량 CPU 모델은 상황실 기존 PC 에서도 동작 → 도입 비용 낮음",
        "유의점: 오분류 위험이 있으므로 대원의 판단을 대체하지 않는 '보조 도구'로 운영해야 함",
    ])
    prs.save(path)
    return path
